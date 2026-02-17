from fastapi import FastAPI, File, Request, UploadFile
from fastapi.responses import HTMLResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from io import BytesIO

app = FastAPI(title="Visor de Excel")
templates = Jinja2Templates(directory="templates")


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse(
        "index.html",
        {
            "request": request,
            "headers": [],
            "rows": [],
            "error": None,
            "filename": None,
        },
    )


@app.post("/upload", response_class=HTMLResponse)
async def upload_excel(request: Request, file: UploadFile = File(...)):
    headers: list[str] = []
    rows: list[list[str]] = []
    error = None

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        error = "Solo se permiten archivos con extensión .xlsx"
        return templates.TemplateResponse(
            "index.html",
            {
                "request": request,
                "headers": headers,
                "rows": rows,
                "error": error,
                "filename": file.filename,
            },
        )

    try:
        content = await file.read()
        workbook = load_workbook(filename=BytesIO(content), data_only=True)
        sheet = workbook.active

        all_rows = list(sheet.iter_rows(values_only=True))

        if not all_rows:
            error = "El archivo está vacío."
        else:
            headers = [str(value) if value is not None else "" for value in all_rows[0]]
            rows = [
                [str(value) if value is not None else "" for value in row]
                for row in all_rows[1:]
            ]
    except Exception:
        error = "No se pudo leer el archivo. Verifica que sea un Excel válido."

    return templates.TemplateResponse(
        "index.html",
        {
            "request": request,
            "headers": headers,
            "rows": rows,
            "error": error,
            "filename": file.filename,
        },
    )
